import docx

def create_sample_docx():
    doc = docx.Document()
    doc.add_heading('Sarah Connor', 0)
    
    p_contact = doc.add_paragraph('Email: sarah.connor@example.com | Phone: +1 555-839-2019 | Location: Austin, TX\nLinkedIn: linkedin.com/in/sarah-connor | GitHub: github.com/sarahconnor-ai')
    
    doc.add_heading('Professional Summary', level=1)
    doc.add_paragraph('Dynamic Data Scientist & Machine Learning Specialist with 2+ years experience building predictive models, NLP text classifiers, and end-to-end data analytics pipelines in Python. Proficient in Scikit-learn, TensorFlow, Pandas, and SQL.')

    doc.add_heading('Technical Skills', level=1)
    doc.add_paragraph('• Programming: Python, R, SQL, Bash\n• ML & AI: Machine Learning, Deep Learning, TensorFlow, Scikit-learn, Pandas, NumPy, NLP, Data Analysis, Data Visualization\n• Databases & Cloud: PostgreSQL, MySQL, AWS, Docker, Git\n• Soft Skills: Problem Solving, Critical Thinking, Communication')

    doc.add_heading('Professional Experience', level=1)
    p_exp1 = doc.add_paragraph('Junior Data Scientist | DataSphere Analytics (2023 - Present)')
    p_exp1.bold = True
    doc.add_paragraph('• Developed predictive customer churn models using Random Forest and XGBoost, increasing retention by 22% across 50,000 customers.\n• Engineered automated data preprocessing and feature extraction pipelines with Pandas and NumPy.\n• Conducted A/B testing on recommendation engine algorithms to optimize click-through conversion rates.')

    doc.add_heading('Education', level=1)
    doc.add_paragraph('Bachelor of Science (B.Sc) in Computer Science & Statistics\nUniversity of Texas at Austin | 2019 - 2023 | GPA: 3.75/4.0')

    doc.add_heading('Key Projects', level=1)
    doc.add_paragraph('• Healthcare Sentiment Analyzer: Fine-tuned NLP text classification model achieving 91% accuracy on clinical feedback.\n• Real-Time Stock Predictor: Built time-series forecasting model using Python and deployed with Streamlit.')

    doc.add_heading('Certifications', level=1)
    doc.add_paragraph('• DeepLearning.AI TensorFlow Developer Specialization\n• AWS Certified Cloud Practitioner')

    doc.save('sample_data_scientist_resume.docx')
    print("✓ Created sample_data_scientist_resume.docx successfully!")

if __name__ == '__main__':
    create_sample_docx()
