import os
import re
import io
import docx
from PyPDF2 import PdfReader

def extract_text_from_pdf(file_obj_or_path):
    """
    Extracts text and page count from a PDF file using PyPDF2.
    Accepts a filepath string or an in-memory uploaded file object.
    """
    text_content = []
    page_count = 0
    
    try:
        if isinstance(file_obj_or_path, str):
            reader = PdfReader(file_obj_or_path)
        else:
            # Read from uploaded file stream
            if hasattr(file_obj_or_path, 'read'):
                file_bytes = file_obj_or_path.read()
                # reset pointer for future operations
                if hasattr(file_obj_or_path, 'seek'):
                    file_obj_or_path.seek(0)
                reader = PdfReader(io.BytesIO(file_bytes))
            else:
                reader = PdfReader(file_obj_or_path)
                
        page_count = len(reader.pages)
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_content.append(page_text)
                
    except Exception as e:
        print(f"Error extracting PDF text: {e}")
        
    full_text = "\n".join(text_content).strip()
    return full_text, page_count


def extract_text_from_docx(file_obj_or_path):
    """
    Extracts text, paragraphs, and tables from a DOCX file using python-docx.
    """
    text_content = []
    
    try:
        if isinstance(file_obj_or_path, str):
            doc = docx.Document(file_obj_or_path)
        else:
            if hasattr(file_obj_or_path, 'read'):
                file_bytes = file_obj_or_path.read()
                if hasattr(file_obj_or_path, 'seek'):
                    file_obj_or_path.seek(0)
                doc = docx.Document(io.BytesIO(file_bytes))
            else:
                doc = docx.Document(file_obj_or_path)
                
        for para in doc.paragraphs:
            if para.text.strip():
                text_content.append(para.text.strip())
                
        # Also extract table text (often used in resumes for education/skills)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_content.append(" | ".join(row_text))
                    
    except Exception as e:
        print(f"Error extracting DOCX text: {e}")
        
    full_text = "\n".join(text_content).strip()
    # Estimate pages for docx: approximately 350-400 words per page
    words = len(full_text.split())
    estimated_pages = max(1, (words + 300) // 350)
    return full_text, estimated_pages


def extract_text_from_file(file_obj_or_path, filename=None):
    """
    Unified text extraction dispatcher based on file extension.
    Returns: dict with text, page_count, word_count, char_count, file_type.
    """
    if filename is None:
        if isinstance(file_obj_or_path, str):
            filename = os.path.basename(file_obj_or_path)
        elif hasattr(file_obj_or_path, 'name'):
            filename = file_obj_or_path.name
        else:
            filename = "document.pdf"
            
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == '.pdf':
        text, pages = extract_text_from_pdf(file_obj_or_path)
    elif ext in ['.docx', '.doc']:
        text, pages = extract_text_from_docx(file_obj_or_path)
    else:
        text, pages = "", 1

    # Clean whitespace and normalize line breaks
    cleaned_text = re.sub(r'[ \t]+', ' ', text)
    cleaned_text = re.sub(r'\n{3,}', '\n\n', cleaned_text)
    
    words = cleaned_text.split()
    word_count = len(words)
    char_count = len(cleaned_text)
    
    return {
        'text': cleaned_text,
        'page_count': max(1, pages),
        'word_count': word_count,
        'char_count': char_count,
        'file_type': ext,
        'file_name': filename,
    }
